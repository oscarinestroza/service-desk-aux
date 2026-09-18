"""Descarga del reporte de tickets del SIG (Excel) y pipeline de tickets.

Fase 1: solo la descarga del Excel. El parser a BD llega en la Fase 2.

Dos modos de descarga (misma función, pasos opcionales):
  - MODO_RAPIDA:   pasos 1 y 5 (ir a la URL + apretar descargar). Sin filtros.
                   Timeout de descarga: 3 minutos.
  - MODO_COMPLETA: pasos 1 a 5 (además activa filtros avanzados, llena la
                   fecha 'desde' y aplica). Timeout de descarga: 10 minutos.

Las credenciales del worker de tickets se leen del entorno
(SIG_TICKETS_URL / SIG_TICKETS_USUARIO / SIG_TICKETS_PASSWORD), a diferencia
de la cuenta de enlaces que vive en ConfiguracionSIG (admin).
"""

import json
import logging
import os
import time
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from celery import shared_task
from django.conf import settings
from django.db import transaction
from django.utils import timezone
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

from ..models import ConfiguracionTickets
from .client import SIGClient, SIGClientError

logger = logging.getLogger("sig_sync_tickets")

MODO_RAPIDA = "rapida"
MODO_COMPLETA = "completa"
MODO_VALORES = (MODO_RAPIDA, MODO_COMPLETA)

# Timeouts de descarga (segundos)
TIMEOUT_DESCARGA_RAPIDA = 3 * 60     # aborta a los 3 min
TIMEOUT_DESCARGA_COMPLETA = 10 * 60  # aborta a los 10 min

REPORTE_URL_DEFAULT = "/admin/Solicitud/SeguimientoAtencion"

# Selectores provisionales del reporte (calibrar en la primera prueba real con
# capturas; se evita el hash jss... de MUI cuando es posible).
SEL_TICKETS = {
    "filtros_avanzados_switch": (
        "input.MuiSwitch-input[type='checkbox']",
        "input[type='checkbox'][value='']",
    ),
    "fecha_desde": (
        "div:has(> label:has-text('Inicio')) input",
        "input[aria-label*='desde' i]",
        "input[placeholder*='desde' i]",
    ),
    "aplicar_filtros": (
        "button:has-text('Aplicar')",
        "button:has-text('Buscar')",
        "button:has-text('Filtrar')",
        "button[type='submit']",
    ),
    "descargar_excel": (
        "#btnSolicitudesExcel",
        "button:has-text('EXPORT')",
        "button:has-text('Excel')",
        "a:has-text('EXPORT')",
        "a:has-text('Excel')",
    ),
}

# Timeout (ms) para esperar el botón de descarga: el reporte del SIG puede
# tardar en renderizar; se le da más margen que a los demás selectores.
TIMEOUT_BOTON_DESCARGA_MS = 20000


def _directorio_descargas() -> Path:
    path = Path(settings.MEDIA_ROOT) / "tickets" / "descargas"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _limpiar_descargas():
    """Elimina los temporales de descarga (Excel y JSON de metadatos).

    Se llama tras cada importación para que media/ no se llene.
    """
    try:
        dirp = _directorio_descargas()
    except Exception:  # noqa: BLE001
        return
    for f in dirp.glob("*"):
        if f.is_file():
            try:
                f.unlink()
            except OSError:
                pass


def _registrar_error_log(mensaje):
    """Anexa una línea al archivo tickets_errores.log (para revisión manual)."""
    try:
        path = Path(settings.BASE_DIR) / "tickets_errores.log"
        with path.open("a", encoding="utf-8") as f:
            f.write(f"{datetime.now():%Y-%m-%d %H:%M:%S} | {mensaje}\n")
    except Exception as e:  # noqa: BLE001
        logger.warning("No se pudo escribir tickets_errores.log: %s", e)


def _normalizar_fecha_desde(valor) -> str:
    """Convierte el valor 'desde' a formato dd/mm/aaaa (como lo pide el SIG).

    Acepta '01/01/2020' (dd/mm/aaaa) o '2020-01-01' (ISO).
    """
    valor = (valor or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(valor, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    return "01/01/2020"


def _normalizar_base_url(url) -> str:
    """Deja la URL base del SIG en su forma canónica (termina en /webapp).

    Tolera que llegue completa (https://sig.gia.mx/webapp/seguridad/entrar)
    o ya base (https://sig.gia.mx/webapp) para no depender del formato.
    """
    url = (url or "").strip().rstrip("/")
    idx = url.find("/webapp")
    if idx != -1:
        return url[: idx + len("/webapp")]
    return url


def _credenciales_tickets():
    """Devuelve (url, usuario, password) de la cuenta del worker de tickets.

    Usa el entorno (env vars); la URL cae a ConfiguracionSIG si falta.
    """
    from ..models import ConfiguracionSIG

    url = os.environ.get("SIG_TICKETS_URL", "").strip()
    if not url:
        try:
            url = ConfiguracionSIG.cargar().url_valor()
        except Exception:  # noqa: BLE001
            url = ""
    url = _normalizar_base_url(url)
    usuario = os.environ.get("SIG_TICKETS_USUARIO", "").strip()
    password = os.environ.get("SIG_TICKETS_PASSWORD", "").strip()
    if not usuario or not password:
        raise SIGClientError(
            "Credenciales del worker de tickets no configuradas: "
            "SIG_TICKETS_USUARIO / SIG_TICKETS_PASSWORD"
        )
    return url, usuario, password


def _fecha_desde_config() -> str:
    try:
        return ConfiguracionTickets.cargar().fecha_desde_valor()
    except Exception:  # noqa: BLE001
        return "01/01/2020"


def _url_reporte_config() -> str:
    try:
        return ConfiguracionTickets.cargar().url_reporte_valor()
    except Exception:  # noqa: BLE001
        return REPORTE_URL_DEFAULT


def _tamano_humano(path: Path) -> str:
    size = path.stat().st_size
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024:
            return f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{size:.1f} TB"


def _escribir_metadatos(datos) -> Path:
    """Guarda un JSON con datos del intento de descarga (para debug)."""
    datos = dict(datos)
    datos.setdefault("fecha", datetime.now().isoformat(timespec="seconds"))
    path = _directorio_descargas() / f"descarga_{datetime.now():%Y%m%d_%H%M%S}.json"
    try:
        path.write_text(
            json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except Exception as e:  # noqa: BLE001
        logger.warning("No se pudo escribir metadatos: %s", e)
    return path


class TicketsSIGClient(SIGClient):
    """Cliente Playwright para descargar el reporte de tickets del SIG."""

    def __init__(self, modo=MODO_RAPIDA, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.modo = modo if modo in MODO_VALORES else MODO_RAPIDA
        self.timeout_descarga_ms = (
            TIMEOUT_DESCARGA_RAPIDA if self.modo == MODO_RAPIDA
            else TIMEOUT_DESCARGA_COMPLETA
        ) * 1000

    def descargar_excel(self, report_url=None, desde=None):
        """Pasos 1 y 5 (+ 2, 3, 4 en modo completo). Devuelve dict con la ruta."""
        url = self._url(report_url or REPORTE_URL_DEFAULT)
        inicio = time.monotonic()
        page = self.page
        timeout_ms = self.timeout * 1000

        logger.info("Tickets (%s): navegando a %s", self.modo, url)
        page.goto(url, timeout=timeout_ms, wait_until="domcontentloaded")

        if self.modo == MODO_COMPLETA:
            self._aplicar_filtros_completa(desde or _fecha_desde_config())

        with page.expect_download(timeout=self.timeout_descarga_ms) as dl_info:
            self._disparar_descarga()

        download = dl_info.value
        nombre = download.suggested_filename or f"tickets_{self.modo}.xlsx"
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        destino = _directorio_descargas() / f"{stamp}_{nombre}"
        download.save_as(str(destino))

        segundos = round(time.monotonic() - inicio, 1)
        logger.info(
            "Tickets (%s): descargado %s (%s) en %s s",
            self.modo, nombre, _tamano_humano(destino), segundos,
        )
        return {
            "modo": self.modo,
            "path": str(destino),
            "nombre": nombre,
            "tamano_bytes": destino.stat().st_size,
            "segundos": segundos,
        }

    def _aplicar_filtros_completa(self, desde):
        """Pasos 2, 3 y 4: filtros avanzados + fecha 'desde' + aplicar."""
        if not self._click_primero(SEL_TICKETS["filtros_avanzados_switch"]):
            raise SIGClientError(
                "No se encontró el switch de filtros avanzados (calibrar selector)"
            )
        if not self._fill_primero(SEL_TICKETS["fecha_desde"], desde):
            raise SIGClientError(
                "No se encontró el campo de fecha 'desde' (calibrar selector)"
            )
        if not self._click_primero(SEL_TICKETS["aplicar_filtros"]):
            raise SIGClientError(
                "No se encontró el botón de aplicar filtros (calibrar selector)"
            )
        self.page.wait_for_timeout(1500)
        logger.info("Tickets (completa): filtros aplicados (desde=%s)", desde)

    def _disparar_descarga(self):
        if not self._click_primero(
            SEL_TICKETS["descargar_excel"], timeout_ms=TIMEOUT_BOTON_DESCARGA_MS
        ):
            raise SIGClientError(
                "No se encontró el botón de descarga del Excel (calibrar selector)"
            )

    def _click_primero(self, selectores, timeout_ms=8000):
        for sel in selectores:
            try:
                loc = self.page.locator(sel).first
                loc.wait_for(state="visible", timeout=timeout_ms)
                loc.click(timeout=timeout_ms)
                return True
            except PlaywrightTimeoutError:
                continue
            except Exception:  # noqa: BLE001
                continue
        return False

    def _fill_primero(self, selectores, valor):
        for sel in selectores:
            try:
                loc = self.page.locator(sel).first
                loc.wait_for(state="visible", timeout=8000)
                loc.fill(str(valor))
                return True
            except PlaywrightTimeoutError:
                continue
            except Exception:  # noqa: BLE001
                continue
        return False


def descargar_excel_tickets(modo=MODO_RAPIDA, desde=None, report_url=None):
    """Orquesta la descarga: credenciales -> login -> pasos -> guardado.

    Regresa dict con los datos de la descarga. Si algo falla guarda screenshot
    y metadatos de error, y re-lanza la excepción.
    """
    base_url, usuario, password = _credenciales_tickets()
    desde = _normalizar_fecha_desde(desde or _fecha_desde_config())
    url_reporte = report_url or _url_reporte_config()

    client = None
    try:
        client = TicketsSIGClient(
            modo=modo, base_url=base_url, username=usuario, password=password
        )
        client.login()
        resultado = client.descargar_excel(report_url=url_reporte, desde=desde)
        _escribir_metadatos(resultado)
        return resultado
    except Exception as e:  # noqa: BLE001
        logger.exception("Descarga de tickets (%s) falló", modo)
        _escribir_metadatos({"modo": modo, "error": str(e)})
        if client is not None:
            try:
                client._screenshot(f"tickets_{modo}_error")
            except Exception:  # noqa: BLE001
                pass
        raise
    finally:
        if client is not None:
            try:
                client.logout()
            except Exception:  # noqa: BLE001
                try:
                    client.close()
                except Exception:  # noqa: BLE001
                    pass


@shared_task(queue="programadas")
def descargar_excel_tickets_task(modo=MODO_RAPIDA, desde=None, report_url=None):
    """Tarea Celery del worker de tickets (cola programadas)."""
    resultado = descargar_excel_tickets(
        modo=modo, desde=desde, report_url=report_url
    )
    logger.info("Tarea descargar_excel_tickets_task terminada: %s", resultado)
    return resultado


# ---------------------------------------------------------------------------
# Fase 2: parser del Excel y sincronización a base de datos
# ---------------------------------------------------------------------------
MODO_SYNC_PARCIAL = "parcial"
MODO_SYNC_COMPLETO = "completo"
MODO_SYNC_VALORES = (MODO_SYNC_PARCIAL, MODO_SYNC_COMPLETO)

LOCK_CLAVE = "enlacesccg:tickets:lock"
LOCK_TTL = int(os.environ.get("TICKETS_LOCK_TTL", "600"))

# Mapeo de encabezados del Excel -> campo del modelo Ticket
_MAPA_CAMPOS_TICKETS = {
    "ID Solicitud servicio": "ticket_id",
    "foliosolicitudservicio": "numero",
    "solicitud_tipo": "solicitud_tipo",
    "solicitud_fecha": "fecha",
    "cerro_fecha": "fecha_cierre",
    "solicitud_descripcion": "descripcion",
}


def _parsear_fecha(valor):
    """Convierte fechas del Excel (str o datetime) a datetime aware o None."""
    if valor is None or str(valor).strip() in ("", "None"):
        return None
    if isinstance(valor, datetime):
        fecha = valor
    else:
        texto = str(valor).strip()
        fecha = None
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M:%S.%f",
            "%Y-%m-%d",
            "%d/%m/%Y %H:%M:%S",
            "%d/%m/%Y",
        ):
            try:
                fecha = datetime.strptime(texto, fmt)
                break
            except ValueError:
                continue
    if fecha is None:
        return None
    if timezone.is_naive(fecha):
        fecha = timezone.make_aware(fecha)
    return fecha


def _valor_crudo(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    return str(valor)


def parsear_excel_tickets(ruta, hoja=None):
    """Lee el Excel del SIG y devuelve (datos, errores).

    - Mapea por encabezados (patrón _MAPA_CAMPOS de enlaces).
    - Dedupe por ticket_id dejando la ÚLTIMA fila de cada id.
    - Cada dict trae además 'raw_data' (snapshot de todas las columnas).
    """
    from ..models import ConfiguracionTickets

    if hoja is None:
        hoja = ConfiguracionTickets.cargar().hoja_valor()

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb[hoja] if hoja else wb[wb.sheetnames[0]]

        it = ws.iter_rows(values_only=True)
        encabezados = [str(h or "").strip() for h in next(it, ())]
        mapa_col = {}
        for i, nombre in enumerate(encabezados):
            campo = _MAPA_CAMPOS_TICKETS.get(nombre)
            if campo:
                mapa_col[campo] = i

        por_id = {}
        errores = 0
        for fila in it:
            fila = {encabezados[i]: fila[i] for i in range(min(len(encabezados), len(fila)))}
            tid = str(fila.get("ID Solicitud servicio") or "").strip()
            if not tid:
                errores += 1
                continue
            por_id[tid] = {
                "ticket_id": tid,
                "numero": str(fila.get("foliosolicitudservicio") or "").strip(),
                "solicitud_tipo": str(fila.get("solicitud_tipo") or "").strip(),
                "fecha": _parsear_fecha(fila.get("solicitud_fecha")),
                "fecha_cierre": _parsear_fecha(fila.get("cerro_fecha")),
                "descripcion": str(fila.get("solicitud_descripcion") or "").strip(),
                "raw_data": {k: _valor_crudo(v) for k, v in fila.items()},
            }

        return list(por_id.values()), errores
    finally:
        # Cierra el workbook para liberar el archivo (en Windows queda
        # bloqueado y la limpieza posterior no podría borrarlo).
        wb.close()


def _calcular_numero_display(datos):
    """Asigna numero_display: número base + sufijo -N a los duplicados.

    En el mismo set, la primera aparición (orden ticket_id) conserva el
    número; las siguientes reciben -2, -3... Los vacíos usan el ticket_id.
    """
    for d in datos:
        d["numero_display"] = d["numero"] or d["ticket_id"]

    con_numero = [d for d in datos if d["numero"]]
    con_numero.sort(key=lambda d: (d["numero"], d["ticket_id"]))
    anterior = None
    n_consecutivo = 2
    for d in con_numero:
        if d["numero"] == anterior:
            d["numero_display"] = f"{d['numero']}-{n_consecutivo}"
            n_consecutivo += 1
        else:
            anterior = d["numero"]
            n_consecutivo = 2


def _chunks(secuencia, n=2500):
    """Parte una lista en bloques (evita explotar el límite de parámetros IN)."""
    for i in range(0, len(secuencia), n):
        yield secuencia[i:i + n]


def _normalizar_nombre(valor) -> str:
    """Normaliza texto para empates: sin acentos, espacios colapsados, MAYÚSCULAS.

    Se aplica a AMBOS lados (Excel y directorio) para que " Cesar  Augusto
    Zavala " empate con "César Augusto Zavala".
    """
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.split()).upper()


def cargar_caches_vinculos():
    """Carga en memoria los catálogos/directorio para resolver vínculos (F3)."""
    from ..models import (
        Edificio, EnlaceAutorizado, Falla, Nivel, ResponsableAtencion, Servicio,
    )

    caches = {
        "edificios": {},
        "enlaces": {},
        "servicios": {},
        "fallas": {},
        "niveles": {},
        "responsables": {},
    }

    for ed in Edificio.objects.all():
        caches["edificios"][_normalizar_nombre(ed.nombre)] = ed
        if ed.siglas:
            caches["edificios"].setdefault(_normalizar_nombre(ed.siglas), ed)

    enlaces = sorted(
        EnlaceAutorizado.objects.select_related("institucion").all(),
        key=lambda e: (e.estado != "ACTIVO", e.pk),
    )
    for en in enlaces:
        clave = _normalizar_nombre(en.nombre_sig)
        if clave:
            caches["enlaces"].setdefault(clave, en)
        completo = " ".join(
            x for x in (en.nombres, en.primer_apellido, en.segundo_apellido) if x
        )
        caches["enlaces"].setdefault(_normalizar_nombre(completo), en)

    for s in Servicio.objects.all():
        caches["servicios"][_normalizar_nombre(s.nombre)] = s
    for f in Falla.objects.all():
        caches["fallas"][_normalizar_nombre(f.descripcion)] = f
    for n in Nivel.objects.all():
        caches["niveles"][_normalizar_nombre(n.nombre)] = n
    for r in ResponsableAtencion.objects.all():
        caches["responsables"][_normalizar_nombre(r.nombre)] = r
    return caches


def _resolver_vinculos(raw, caches):
    """Resuelve los vínculos de un ticket a partir de su raw_data (Fase 3)."""
    from ..models import Edificio, Falla, Nivel, ResponsableAtencion, Servicio

    raw = raw or {}
    edificio_nombre = str(raw.get("nivel") or "").strip()
    nivel_nombre = str(raw.get("grupo") or "").strip()
    solicitante_raw = str(raw.get("solicitud_solicitante") or "").strip()
    servicio_nombre = str(raw.get("servicio") or "").strip()
    falla_desc = str(raw.get("falla_descripcion") or "").strip()
    responsable_nombre = str(raw.get("Responsable_atencion") or "").strip()

    # Edificio: si el SIG reporta uno que no existe en el directorio, se crea.
    torre = None
    if edificio_nombre:
        clave_ed = _normalizar_nombre(edificio_nombre)
        torre = caches["edificios"].get(clave_ed)
        if torre is None:
            torre, _ = Edificio.objects.get_or_create(nombre=edificio_nombre)
            caches["edificios"][clave_ed] = torre

    solicitante = caches["enlaces"].get(_normalizar_nombre(solicitante_raw))

    servicio = None
    if servicio_nombre:
        clave = _normalizar_nombre(servicio_nombre)
        servicio = caches["servicios"].get(clave)
        if servicio is None:
            servicio, _ = Servicio.objects.get_or_create(nombre=servicio_nombre)
            caches["servicios"][clave] = servicio

    falla = None
    if falla_desc:
        clave = _normalizar_nombre(falla_desc)
        falla = caches["fallas"].get(clave)
        if falla is None:
            falla, _ = Falla.objects.get_or_create(descripcion=falla_desc)
            caches["fallas"][clave] = falla

    nivel = None
    if nivel_nombre:
        clave = _normalizar_nombre(nivel_nombre)
        nivel = caches["niveles"].get(clave)
        if nivel is None:
            nivel, _ = Nivel.objects.get_or_create(nombre=nivel_nombre)
            caches["niveles"][clave] = nivel

    responsable = None
    if responsable_nombre:
        clave = _normalizar_nombre(responsable_nombre)
        responsable = caches["responsables"].get(clave)
        if responsable is None:
            responsable, _ = ResponsableAtencion.objects.get_or_create(
                nombre=responsable_nombre
            )
            caches["responsables"][clave] = responsable

    return {
        "torre": torre,
        "institucion": solicitante.institucion if solicitante else None,
        "solicitante": solicitante,
        "servicio": servicio,
        "falla": falla,
        "nivel": nivel,
        "responsable_atencion": responsable,
        "solicitante_nombre": "" if solicitante else solicitante_raw,
    }


def aplicar_tickets(datos, modo, ahora=None):
    """Guarda los tickets en BD (transaccional). Devuelve conteos.

    - PARCIAL:   insert/update; NO toca los ausentes del Excel.
    - COMPLETO:  insert/update de todas las filas + archivado=True a los
                 ausentes (nunca borra registros).

    Optimización: los existentes se cargan en una sola consulta (por bloques)
    y solo se reescriben las filas que realmente cambiaron (evita el
    re-update masivo en cada descarga completa de ~150 mil registros).
    """
    from ..models import Ticket

    if modo not in MODO_SYNC_VALORES:
        modo = MODO_SYNC_PARCIAL
    if ahora is None:
        ahora = timezone.now()

    creados = actualizados = 0
    with transaction.atomic():
        datos = list(datos)
        caches = cargar_caches_vinculos()
        presentes = [d["ticket_id"] for d in datos]
        db_ids = set(Ticket.objects.values_list("ticket_id", flat=True))

        existentes = {}
        for chunk in _chunks(list(db_ids & set(presentes))):
            for t in Ticket.objects.filter(ticket_id__in=list(chunk)):
                existentes[t.ticket_id] = t

        for d in datos:
            estatus = (
                Ticket.ESTATUS_CERRADO if d.get("fecha_cierre")
                else Ticket.ESTATUS_ABIERTO
            )
            defaults = {
                "numero": d.get("numero") or "",
                "numero_display": d.get("numero_display") or d["ticket_id"],
                "solicitud_tipo": d.get("solicitud_tipo") or "",
                "fecha": d.get("fecha"),
                "fecha_cierre": d.get("fecha_cierre"),
                "estatus": estatus,
                "descripcion": d.get("descripcion") or "",
                "raw_data": d.get("raw_data") or {},
                "ultima_sync": ahora,
            }
            defaults.update(_resolver_vinculos(defaults["raw_data"], caches))
            if modo == MODO_SYNC_COMPLETO:
                # Al existir de nuevo en el histórico se des-archiva
                defaults["archivado"] = False

            t = existentes.get(d["ticket_id"])
            if t is None:
                Ticket.objects.create(ticket_id=d["ticket_id"], **defaults)
                creados += 1
                continue

            sin_cambios = (
                t.numero == defaults["numero"]
                and t.numero_display == defaults["numero_display"]
                and t.solicitud_tipo == defaults["solicitud_tipo"]
                and t.fecha == defaults["fecha"]
                and t.fecha_cierre == defaults["fecha_cierre"]
                and t.estatus == estatus
                and t.descripcion == defaults["descripcion"]
                and t.raw_data == defaults["raw_data"]
                and t.torre_id == (defaults["torre"].pk if defaults["torre"] else None)
                and t.institucion_id == (
                    defaults["institucion"].pk if defaults["institucion"] else None
                )
                and t.solicitante_id == (
                    defaults["solicitante"].pk if defaults["solicitante"] else None
                )
                and t.servicio_id == (
                    defaults["servicio"].pk if defaults["servicio"] else None
                )
                and t.falla_id == (defaults["falla"].pk if defaults["falla"] else None)
                and t.nivel_id == (defaults["nivel"].pk if defaults["nivel"] else None)
                and t.responsable_atencion_id == (
                    defaults["responsable_atencion"].pk
                    if defaults["responsable_atencion"]
                    else None
                )
                and t.solicitante_nombre == defaults["solicitante_nombre"]
                and (modo != MODO_SYNC_COMPLETO or not t.archivado)
            )
            if sin_cambios:
                continue
            for campo, valor in defaults.items():
                setattr(t, campo, valor)
            t.save(update_fields=list(defaults.keys()))
            actualizados += 1

        archivados = 0
        if modo == MODO_SYNC_COMPLETO and datos:
            ausentes = sorted(db_ids - set(presentes))
            for chunk in _chunks(ausentes):
                archivados += Ticket.objects.filter(
                    ticket_id__in=list(chunk), archivado=False
                ).update(archivado=True)

    return {
        "creados": creados,
        "actualizados": actualizados,
        "archivados": archivados,
    }


def _registrar_resultado(modo, estado, mensaje, conteos=None):
    """Escribe TicketLog y actualiza la info de ConfiguracionTickets.

    `ultimo_intento` se actualiza siempre; `ultima_sincronizacion` solo cuando
    el estado es OK, para que el badge "Última" no se reinicie si falla.
    """
    from ..models import ConfiguracionTickets, Ticket, TicketLog

    conteos = conteos or {}
    cfg = ConfiguracionTickets.cargar()
    TicketLog.objects.create(
        modo=modo,
        estado=estado,
        mensaje=mensaje,
        creados=conteos.get("creados", 0),
        actualizados=conteos.get("actualizados", 0),
        archivados=conteos.get("archivados", 0),
        errores=conteos.get("errores", 0),
    )
    ahora = timezone.now()
    cfg.ultimo_intento = ahora
    if estado == TicketLog.ESTADO_OK:
        cfg.ultima_sincronizacion = ahora
    cfg.ultimo_modo = modo
    cfg.ultimo_estado = estado
    cfg.ultimo_mensaje = mensaje
    cfg.total_tickets = Ticket.objects.count()
    cfg.save(
        update_fields=[
            "ultimo_intento", "ultima_sincronizacion", "ultimo_modo",
            "ultimo_estado", "ultimo_mensaje", "total_tickets",
        ]
    )
    if estado != TicketLog.ESTADO_OK:
        _registrar_error_log(f"[{modo}] {mensaje}")


def sincronizar_tickets(modo=None, desde=None):
    """Descarga el Excel según el modo y lo aplica a la BD.

    Devuelve dict con conteos. El Excel descargado se elimina tras
    importarse para no acumular decenas de MB en media/.
    """
    from ..models import ConfiguracionTickets, TicketLog

    cfg = ConfiguracionTickets.cargar()
    if modo is None:
        modo = cfg.modo_default_valor()
    if modo not in MODO_SYNC_VALORES:
        modo = MODO_SYNC_PARCIAL
    modo_descarga = MODO_COMPLETA if modo == MODO_SYNC_COMPLETO else MODO_RAPIDA

    resultado = descargar_excel_tickets(modo=modo_descarga, desde=desde)
    ruta = resultado["path"]
    try:
        datos, errores_parse = parsear_excel_tickets(ruta, cfg.hoja_valor())
        _calcular_numero_display(datos)
        conteos = aplicar_tickets(datos, modo)
    finally:
        # Borra el Excel descargado y los JSON de metadatos para no llenar media/.
        _limpiar_descargas()

    conteos["errores"] = errores_parse
    conteos["modo"] = modo
    conteos["filas"] = len(datos)
    mensaje = (
        f"{modo}: {len(datos)} filas, {conteos['creados']} creados, "
        f"{conteos['actualizados']} actualizados, {conteos['archivados']} archivados."
        + (f" {errores_parse} filas con error." if errores_parse else "")
    )
    _registrar_resultado(modo, TicketLog.ESTADO_OK, mensaje, conteos)
    return conteos


def _adquirir_lock():
    """Lock Redis (set nx ex). False = ya ocupado; None = sin Redis (procede)."""
    import redis  # noqa: PLC0415

    try:
        r = redis.from_url(settings.CELERY_BROKER_URL)
        ok = r.set(LOCK_CLAVE, "1", nx=True, ex=LOCK_TTL)
    except Exception as e:  # noqa: BLE001
        logger.warning("Redis no disponible para lock; se procede sin lock: %s", e)
        return None
    return r if ok else False


def _liberar_lock(redis_client):
    if redis_client:
        try:
            redis_client.delete(LOCK_CLAVE)
        except Exception:  # noqa: BLE001
            pass


def _lock_activo():
    """True si hay una sincronización en curso (lock Redis). None si sin Redis."""
    import redis  # noqa: PLC0415

    try:
        r = redis.from_url(settings.CELERY_BROKER_URL)
        return bool(r.exists(LOCK_CLAVE))
    except Exception as e:  # noqa: BLE001
        logger.warning("Redis no disponible para consultar lock: %s", e)
        return None


@shared_task(queue="programadas")
def sincronizar_tickets_task(modo=None, desde=None):
    """Tarea para 'Sincronizar ahora' (botón). Respeta el lock Redis."""
    lock = _adquirir_lock()
    if lock is False:
        logger.info("Sincronización omitida: ya hay una corrida en curso")
        return {"sincronizado": False, "motivo": "lock"}
    try:
        resultado = sincronizar_tickets(modo=modo, desde=desde)
        return {"sincronizado": True, **resultado}
    finally:
        _liberar_lock(lock)


@shared_task(queue="programadas")
def tick_sync_tickets_task():
    """Tick de 60 s: decide modo y sincroniza si toca.

    - Deshabilitado -> no corre.
    - Hora en descarga_completa_horas -> completo.
    - Si no, parcial cuando pasó intervalo_minutos desde el último sync.
    """
    from ..models import ConfiguracionTickets, TicketLog

    cfg = ConfiguracionTickets.cargar()
    if not cfg.habilitado:
        return {"sincronizado": False, "motivo": "deshabilitado"}

    ahora = timezone.now()
    horas = cfg.descarga_completa_horas or []
    # Las horas configuradas son locales; comparar en hora local (no UTC).
    hora_local = timezone.localtime(ahora).strftime("%H:%M")
    es_hora_completa = isinstance(horas, list) and hora_local in horas
    if es_hora_completa:
        modo = MODO_SYNC_COMPLETO
    else:
        # La sincronización parcial solo corre dentro del horario/días
        # configurados. La completa (arriba) no depende del horario, y el
        # botón "Sincronizar ahora" tampoco (usa sincronizar_tickets directo).
        if not cfg.en_horario(ahora):
            return {"sincronizado": False, "motivo": "fuera_horario"}
        # El intervalo se mide desde el último INTENTO (aunque haya fallado),
        # para no reintentar en cada tick de 60 s tras un error.
        referencia = cfg.ultimo_intento or cfg.ultima_sincronizacion
        intervalo = cfg.intervalo_minutos or 5
        if referencia and (ahora - referencia) < timedelta(minutes=intervalo):
            return {"sincronizado": False, "motivo": "intervalo"}
        modo = MODO_SYNC_PARCIAL

    lock = _adquirir_lock()
    if lock is False:
        return {"sincronizado": False, "motivo": "lock"}
    try:
        resultado = sincronizar_tickets(modo=modo)
        return {"sincronizado": True, **resultado}
    except Exception as e:  # noqa: BLE001
        logger.exception("Tick de sincronización de tickets falló (%s)", modo)
        _registrar_resultado(modo, TicketLog.ESTADO_ERROR, str(e))
        return {"sincronizado": False, "motivo": "error", "error": str(e)}
    finally:
        _liberar_lock(lock)