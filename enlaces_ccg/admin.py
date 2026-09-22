"""
Registro de modelos en el sitio de administración de Django.
"""

from io import BytesIO
from datetime import datetime

from django.contrib import admin, messages
from django import forms
from django.http import HttpResponse
from django.template.response import TemplateResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from .models import (
    Adjunto, ComunicadoCC, ConfiguracionCorreo, ConfiguracionSIG,
    ConfiguracionTickets, Documento, DocumentoCarpeta,
    Edificio,
    EnlaceAutorizado,
    Falla,
    Institucion,
    InstitucionEdificio,
    Nivel,
    ResponsableAtencion,
    Seccion,
    Servicio,
    SyncLog,
    Ticket,
    TicketAdjunto,
    TicketCierre,
    TicketLog,
    TicketRegistro,
    VistaGuardada,
)


# ---------------------------------------------------------------------------
# Estilos Excel reutilizables
# ---------------------------------------------------------------------------
_HEADER_FILL = PatternFill(start_color="0D1B3E", end_color="0D1B3E", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_DATA_FONT = Font(name="Calibri", size=10)
_TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="0D1B3E")
_SUBTITLE_FONT = Font(name="Calibri", bold=False, size=10, color="666666")
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _generar_excel_enlaces(enlaces, titulo, subtitulo, archivo_prefijo):
    """Genera un HttpResponse con un archivo Excel de enlaces."""
    COLUMNAS = [
        "Nombres", "Pr. Apellido", "Seg. Apellido", "Genero",
        "Correo Principal", "Correo Secundario",
        "Telefono 1", "Telefono 2", "Extension",
        "Estado", "Nivel Ref.", "Usuario SIG", "Nombre SIG",
        "Institucion", "Siglas", "Edificio",
        "Fecha Alta", "Oficio Alta", "Observ. Alta",
        "Fecha Seg.", "Oficio Seg.", "Observ. Seg.",
        "Fecha Baja", "Oficio Baja", "Observ. Baja",
    ]

    def _fila(e):
        return (
            e.nombres, e.primer_apellido, e.segundo_apellido or "",
            e.get_genero_display() if e.genero else "",
            e.correo_principal, e.correo_secundario or "",
            e.telefono_principal, e.telefono_secundario or "",
            e.extension_telefonica or "", e.estado,
            e.nivel_referencia or "", e.usuario_sig or "", e.nombre_sig or "",
            e.institucion.nombre if e.institucion else "",
            e.institucion.siglas or "" if e.institucion else "",
            ", ".join(ed.nombre for ed in e.institucion.edificio.all()) if e.institucion else "",
            e.fecha_alta.strftime("%d/%m/%Y") if e.fecha_alta else "",
            e.oficio_alta or "", e.observaciones_alta or "",
            e.fecha_seguimiento.strftime("%d/%m/%Y") if e.fecha_seguimiento else "",
            e.oficio_seguimiento or "", e.observaciones_seguimiento or "",
            e.fecha_baja.strftime("%d/%m/%Y") if e.fecha_baja else "",
            e.oficio_baja or "", e.observaciones_baja or "",
        )

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    num_col = len(COLUMNAS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_col)
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = _TITLE_FONT
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_col)
    c = ws.cell(row=2, column=1, value=subtitulo)
    c.font = _SUBTITLE_FONT
    ws.row_dimensions[2].height = 20

    for ci, nombre in enumerate(COLUMNAS, 1):
        c = ws.cell(row=3, column=ci, value=nombre)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _THIN_BORDER

    for ri, enlace in enumerate(enlaces, 4):
        for ci, val in enumerate(_fila(enlace), 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font = _DATA_FONT
            c.border = _THIN_BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)

    for ci in range(1, num_col + 1):
        max_len = len(str(COLUMNAS[ci - 1]))
        for ri in range(4, ws.max_row + 1):
            val = ws.cell(row=ri, column=ci).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[ws.cell(row=3, column=ci).column_letter].width = min(max_len + 4, 40)

    total = ws.max_row + 2
    ws.cell(row=total, column=1, value=f"Total: {enlaces.count()} registros").font = Font(
        name="Calibri", bold=True, size=10, color="333333"
    )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{archivo_prefijo}_{ts}.xlsx"'
    return resp


# ---------------------------------------------------------------------------
# Edificio
# ---------------------------------------------------------------------------
@admin.register(Edificio)
class EdificioAdmin(admin.ModelAdmin):
    list_display = ("nombre", "siglas", "visible", "creado_en")
    list_editable = ("visible",)
    list_filter = ("visible",)
    search_fields = ("nombre", "siglas")
    readonly_fields = ("creado_en",)

    fieldsets = (
        (None, {
            "fields": ("nombre", "siglas", "descripcion", "visible"),
        }),
        ("Metadata", {
            "classes": ("collapse",),
            "fields": ("creado_en",),
        }),
    )


# ---------------------------------------------------------------------------
# ComunicadoCC
# ---------------------------------------------------------------------------
@admin.register(ComunicadoCC)
class ComunicadoCCAdmin(admin.ModelAdmin):
    list_display = ("nombre", "institucion", "cargo", "correo", "actualizado_en")
    search_fields = ("nombre", "cargo", "correo", "institucion__nombre")
    list_select_related = ("institucion",)
    autocomplete_fields = ("institucion",)


# ---------------------------------------------------------------------------
# InstitucionEdificio (inline en InstitucionAdmin)
# ---------------------------------------------------------------------------
class InstitucionEdificioInline(admin.TabularInline):
    model = InstitucionEdificio
    extra = 1
    fields = ("edificio", "nivel")
    autocomplete_fields = ("edificio",)


# ---------------------------------------------------------------------------
# Institucion
# ---------------------------------------------------------------------------
@admin.register(Institucion)
class InstitucionAdmin(admin.ModelAdmin):
    change_list_template = "admin/enlaces_ccg/institucion/change_list.html"

    list_display = ("nombre", "siglas", "edificios_display", "estado", "creado_en")
    list_filter = ("edificio", "estado")
    search_fields = ("nombre", "siglas")

    inlines = [InstitucionEdificioInline]

    fieldsets = (
        (None, {
            "fields": ("nombre", "siglas", "estado"),
        }),
        ("Contacto", {
            "classes": ("collapse",),
            "fields": ("contacto_nombre", "contacto_correo", "contacto_telefono"),
        }),
    )

    @admin.display(description="Edificio(s)")
    def edificios_display(self, obj):
        edificios = obj.edificio.all()
        if not edificios:
            return "—"
        return ", ".join(ed.nombre for ed in edificios)


# ---------------------------------------------------------------------------
# EnlaceAutorizado
# ---------------------------------------------------------------------------
@admin.register(EnlaceAutorizado)
class EnlaceAutorizadoAdmin(admin.ModelAdmin):
    change_list_template = "admin/enlaces_ccg/enlaceautorizado/change_list.html"

    list_display = (
        "nombres",
        "primer_apellido",
        "segundo_apellido",
        "institucion",
        "estado",
        "sincronizado",
        "correo_principal",
        "usuario_sig",
    )
    list_filter = ("institucion", "estado", "genero")
    search_fields = (
        "nombres", "primer_apellido", "segundo_apellido",
        "correo_principal", "usuario_sig", "nombre_sig",
    )
    list_select_related = ("institucion",)
    list_editable = ("estado",)
    readonly_fields = ("creado_en", "actualizado_en", "nombre_completo_admin", "edificio")

    fieldsets = (
        ("Datos Generales", {
            "fields": (
                "nombre_completo_admin", "institucion", "edificio",
                "estado", "nivel_referencia",
            ),
        }),
        ("Datos Personales", {
            "fields": ("nombres", "primer_apellido", "segundo_apellido"),
        }),
        ("Contacto", {
            "fields": (
                "correo_principal", "correo_secundario",
                "telefono_principal", "telefono_secundario", "extension_telefonica",
            ),
        }),
        ("Sistema SIG", {
            "classes": ("collapse",),
            "fields": (
                "genero",
                "usuario_sig", "password_sig", "nombre_sig", "pin_sig",
            ),
        }),
        ("Alta, Seguimiento y Baja", {
            "fields": (
                "fecha_alta", "oficio_alta", "observaciones_alta",
                "fecha_seguimiento", "oficio_seguimiento", "observaciones_seguimiento",
                "fecha_baja", "oficio_baja", "observaciones_baja",
            ),
        }),
        ("Comentarios", {
            "classes": ("collapse",),
            "fields": ("comentarios",),
        }),
        ("Timestamps", {
            "classes": ("collapse",),
            "fields": ("creado_en", "actualizado_en"),
        }),
    )

    # -- Acciones batch desde la lista --
    actions = ("exportar_seleccionados_excel",)

    @admin.action(description="Exportar seleccionados a Excel (.xlsx)")
    def exportar_seleccionados_excel(self, request, queryset):
        enlaces = queryset.select_related("institucion").prefetch_related(
            "institucion__edificio"
        ).order_by(
            "institucion__nombre", "primer_apellido", "segundo_apellido", "nombres"
        )
        sub = f"Generado {datetime.now().strftime('%d/%m/%Y %H:%M')} — {request.user.username}"
        return _generar_excel_enlaces(enlaces, "Enlaces Seleccionados", sub, "enlaces_seleccion")

    # -- Vista custom del change_list para agregar botón de exportar todo --
    def changelist_view(self, request, extra_context=None):
        if request.POST.get("action") == "exportar_todos_excel":
            enlaces = EnlaceAutorizado.objects.select_related(
                "institucion"
            ).prefetch_related(
                "institucion__edificio"
            ).order_by("institucion__nombre", "primer_apellido", "segundo_apellido", "nombres")
            sub = f"Todos los enlaces — Generado {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            return _generar_excel_enlaces(enlaces, "Todos los Enlaces CCG", sub, "enlaces_todos")

        return super().changelist_view(request, extra_context=extra_context)

    # -- Exportar un solo enlace desde el formulario de edición --
    def changeform_view(self, request, object_id=None, form_url="", extra_context=None):
        extra_context = extra_context or {}

        if object_id:
            enlace = self.get_object(request, int(object_id))
            if enlace:
                extra_context["enlace_id"] = object_id

        return super().changeform_view(request, object_id, form_url, extra_context)

    def response_change(self, request, obj):
        if "_exportar_excel" in request.POST:
            enlaces = EnlaceAutorizado.objects.filter(pk=obj.pk).select_related(
                "institucion"
            ).prefetch_related("institucion__edificio")
            sub = f"{obj.nombre_completo} — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            return _generar_excel_enlaces(enlaces, f"Enlace: {obj.nombre_completo}", sub, f"enlace_{obj.pk}")

        return super().response_change(request, obj)


# ---------------------------------------------------------------------------
# Adjunto
# ---------------------------------------------------------------------------
@admin.register(Adjunto)
class AdjuntoAdmin(admin.ModelAdmin):
    list_display = ("nombre", "enlace", "subido_en")
    search_fields = ("nombre", "enlace__nombres", "enlace__primer_apellido")
    list_select_related = ("enlace",)


# ---------------------------------------------------------------------------
# SyncLog
# ---------------------------------------------------------------------------
@admin.register(SyncLog)
class SyncLogAdmin(admin.ModelAdmin):
    list_display = (
        "enlace", "accion", "estado", "usuario_sig_id",
        "mensaje_corto", "creado_en",
    )
    list_filter = ("accion", "estado")
    search_fields = (
        "enlace__nombres", "enlace__primer_apellido",
        "enlace__usuario_sig", "usuario_sig_id",
    )
    list_select_related = ("enlace",)
    readonly_fields = ("enlace", "accion", "estado", "mensaje", "usuario_sig_id", "creado_en")

    @admin.display(description="Mensaje")
    def mensaje_corto(self, obj):
        msg = obj.mensaje or ""
        return msg[:80] + "…" if len(msg) > 80 else msg


# ---------------------------------------------------------------------------
# Configuración de correo (singleton)
# ---------------------------------------------------------------------------
class ConfiguracionCorreoForm(forms.ModelForm):
    class Meta:
        model = ConfiguracionCorreo
        fields = "__all__"
        widgets = {
            "password_emisor": forms.PasswordInput(
                render_value=True,
                attrs={"autocomplete": "new-password"},
            ),
        }

    def clean(self):
        cleaned = super().clean()
        # Validar que si hay correo_emisor, también haya password
        emisor = cleaned.get("correo_emisor")
        password = cleaned.get("password_emisor")
        if emisor and not password:
            self.add_error("password_emisor", "Requerido si se define un correo emisor.")
        # Validar puerto SMTP
        port = cleaned.get("smtp_port")
        if port is not None and (port < 1 or port > 65535):
            self.add_error("smtp_port", "Puerto inválido (1-65535).")
        return cleaned


@admin.register(ConfiguracionCorreo)
class ConfiguracionCorreoAdmin(admin.ModelAdmin):
    form = ConfiguracionCorreoForm
    list_display = (
        "correo_emisor", "smtp_host", "smtp_port", "smtp_use_tls",
        "correo_review", "correo_notificacion", "actualizado_en",
    )
    fieldsets = (
        ("Servidor SMTP", {
            "fields": ("smtp_host", "smtp_port", "smtp_use_tls"),
            "description": "Configuración del servidor de correo saliente.",
        }),
        ("Cuenta emisora", {
            "fields": ("correo_emisor", "password_emisor", "default_from_email"),
            "description": "Credenciales de la cuenta que enviará los correos.",
        }),
        ("Destinatarios", {
            "fields": ("correo_review", "correo_notificacion"),
            "description": "Listas separadas por punto y coma (;).",
        }),
    )
    readonly_fields = ("actualizado_en", "creado_en")

    def has_add_permission(self, request):
        # Permitir solo un registro singleton
        return not self.model.objects.exists()

    def has_delete_permission(self, request, obj=None):
        return False

    def changeform_view(self, request, object_id=None, form_url="", extra_context=None):
        # Garantiza que exista el registro singleton antes de abrir el form
        obj, _ = ConfiguracionCorreo.objects.get_or_create(pk=1)
        return super().changeform_view(
            request, str(obj.pk), form_url, extra_context
        )


class ConfiguracionSIGForm(forms.ModelForm):
    class Meta:
        model = ConfiguracionSIG
        fields = "__all__"
        widgets = {
            "password": forms.PasswordInput(
                render_value=True,
                attrs={"autocomplete": "new-password"},
            ),
            "default_password": forms.PasswordInput(
                render_value=True,
                attrs={"autocomplete": "new-password"},
            ),
        }

    def clean(self):
        cleaned = super().clean()
        # Validar que si hay usuario, también haya password
        usuario = cleaned.get("usuario")
        password = cleaned.get("password")
        if usuario and not password:
            self.add_error("password", "Requerido si se define un usuario.")
        return cleaned


@admin.register(ConfiguracionSIG)
class ConfiguracionSIGAdmin(admin.ModelAdmin):
    form = ConfiguracionSIGForm
    list_display = ("url", "usuario", "timeout", "actualizado_en")
    fieldsets = (
        ("Conexión SIG", {
            "fields": ("url", "usuario", "password"),
            "description": "Credenciales para el login automatizado en el SIG.",
        }),
        ("Opciones", {
            "fields": ("default_password", "timeout"),
            "description": "Contraseña por defecto para usuarios creados y timeout de Playwright.",
        }),
    )
    readonly_fields = ("actualizado_en", "creado_en")

    def has_add_permission(self, request):
        # Permitir solo un registro singleton
        return not self.model.objects.exists()

    def has_delete_permission(self, request, obj=None):
        return False

    def changeform_view(self, request, object_id=None, form_url="", extra_context=None):
        # Garantiza que exista el registro singleton antes de abrir el form
        obj, _ = ConfiguracionSIG.objects.get_or_create(pk=1)
        return super().changeform_view(
            request, str(obj.pk), form_url, extra_context
        )


DIAS_SEMANA_CHOICES = [
    (0, "Lunes"), (1, "Martes"), (2, "Miércoles"), (3, "Jueves"),
    (4, "Viernes"), (5, "Sábado"), (6, "Domingo"),
]


class ConfiguracionTicketsForm(forms.ModelForm):
    """Form del singleton de tickets: `dias_semana` como checkboxes."""

    dias_semana = forms.MultipleChoiceField(
        choices=DIAS_SEMANA_CHOICES,
        required=False,
        widget=forms.CheckboxSelectMultiple,
        label="Días permitidos (parcial)",
        help_text="Días en que corre la sincronización parcial. Vacío = todos.",
    )

    class Meta:
        model = ConfiguracionTickets
        fields = "__all__"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        valor = getattr(self.instance, "dias_semana", None) or []
        self.initial["dias_semana"] = [str(x) for x in valor]

    def clean_dias_semana(self):
        return [int(x) for x in self.cleaned_data.get("dias_semana", [])]


@admin.register(ConfiguracionTickets)
class ConfiguracionTicketsAdmin(admin.ModelAdmin):
    form = ConfiguracionTicketsForm
    list_display = (
        "habilitado", "modo_default", "intervalo_minutos", "url_reporte",
        "total_tickets", "ultima_sincronizacion",
    )
    fieldsets = (
        ("Reporte de tickets del SIG", {
            "fields": ("url_reporte", "hoja_excel"),
            "description": (
                "URL del reporte de seguimiento de solicitudes y hoja del Excel "
                "(vacío = auto-detectar la primera)."
            ),
        }),
        ("Descarga completa", {
            "fields": ("fecha_desde_completa",),
            "description": (
                "Fecha 'desde' usada solo en la descarga completa (filtro de "
                "tiempo en el SIG; 'hasta' = hoy). La descarga rápida no usa filtros."
            ),
        }),
        ("Sincronización automática", {
            "fields": (
                "habilitado", "modo_default", "intervalo_minutos",
                "hora_inicio", "hora_fin", "dias_semana",
                "descarga_completa_horas",
            ),
            "description": (
                "La sincronización PARCIAL (cada intervalo_minutos) solo corre "
                "dentro del horario y días indicados (hora local). Vacío = sin "
                "límite. La descarga COMPLETA usa sus horas configuradas y no "
                "depende del horario. El botón 'Sincronizar ahora' funciona siempre."
            ),
        }),
        ("Última sincronización", {
            "fields": (
                "ultima_sincronizacion", "ultimo_modo", "ultimo_estado",
                "ultimo_mensaje", "total_tickets",
            ),
            "description": "Información de la última corrida (solo lectura).",
        }),
    )
    readonly_fields = (
        "ultima_sincronizacion", "ultimo_modo", "ultimo_estado",
        "ultimo_mensaje", "total_tickets", "actualizado_en", "creado_en",
    )

    def has_add_permission(self, request):
        # Permitir solo un registro singleton
        return not self.model.objects.exists()

    def has_delete_permission(self, request, obj=None):
        return False

    def changeform_view(self, request, object_id=None, form_url="", extra_context=None):
        # Garantiza que exista el registro singleton antes de abrir el form
        obj, _ = ConfiguracionTickets.objects.get_or_create(pk=1)
        return super().changeform_view(
            request, str(obj.pk), form_url, extra_context
        )


@admin.register(Servicio)
class ServicioAdmin(admin.ModelAdmin):
    list_display = ("nombre", "activo", "creado_en")
    list_filter = ("activo",)
    search_fields = ("nombre",)
    filter_horizontal = ("responsables",)


@admin.register(Falla)
class FallaAdmin(admin.ModelAdmin):
    list_display = ("descripcion", "activo", "creado_en")
    list_filter = ("activo",)
    search_fields = ("descripcion",)


@admin.register(Nivel)
class NivelAdmin(admin.ModelAdmin):
    list_display = ("nombre", "activo", "creado_en")
    list_filter = ("activo",)
    search_fields = ("nombre",)


@admin.register(ResponsableAtencion)
class ResponsableAtencionAdmin(admin.ModelAdmin):
    list_display = ("nombre", "activo", "creado_en")
    list_filter = ("activo",)
    search_fields = ("nombre",)
    filter_horizontal = ("usuarios",)
    readonly_fields = ("nombre",)


@admin.register(VistaGuardada)
class VistaGuardadaAdmin(admin.ModelAdmin):
    list_display = ("nombre", "usuario", "modulo", "es_predeterminada", "creada_en")
    list_filter = ("modulo", "es_predeterminada")
    search_fields = ("nombre", "usuario__username")


@admin.register(Ticket)
class TicketAdmin(admin.ModelAdmin):
    list_display = (
        "numero_display", "fecha", "estatus", "solicitud_tipo",
        "torre", "nivel", "servicio", "responsable_atencion", "archivado",
        "ultima_sync",
    )
    list_filter = (
        "estatus", "archivado", "solicitud_tipo", "torre", "nivel", "servicio",
        "responsable_atencion",
    )
    search_fields = ("numero_display", "numero", "ticket_id", "descripcion")
    list_select_related = ("torre", "nivel", "servicio", "responsable_atencion")
    ordering = ("-fecha",)
    readonly_fields = (
        "numero", "ticket_id", "fecha", "fecha_cierre", "estatus",
        "solicitud_tipo", "archivado", "ultima_sync", "numero_display",
        "descripcion", "raw_data", "creado_en", "actualizado_en",
        "torre", "institucion", "solicitante", "servicio", "falla",
        "nivel", "responsable_atencion", "solicitante_nombre",
    )

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False


@admin.register(TicketLog)
class TicketLogAdmin(admin.ModelAdmin):
    list_display = ("creado_en", "modo", "estado", "creados", "actualizados", "archivados", "errores")
    list_filter = ("modo", "estado")
    readonly_fields = ("creado_en",)
    date_hierarchy = "creado_en"

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False


@admin.register(TicketRegistro)
class TicketRegistroAdmin(admin.ModelAdmin):
    list_display = ("creado_en", "ticket", "usuario", "descripcion")
    list_filter = ("creado_en",)
    search_fields = ("ticket__numero_display", "ticket__numero", "descripcion")
    readonly_fields = ("creado_en",)
    date_hierarchy = "creado_en"

    def has_add_permission(self, request):
        return False


@admin.register(TicketCierre)
class TicketCierreAdmin(admin.ModelAdmin):
    list_display = (
        "ticket",
        "fecha_inicio",
        "fecha_cierre",
        "actualizado_por",
        "actualizado_en",
    )
    list_filter = ("actualizado_en",)
    search_fields = ("ticket__numero_display", "ticket__numero", "ticket__ticket_id")
    readonly_fields = ("creado_en", "actualizado_en", "creado_por", "actualizado_por")


@admin.register(TicketAdjunto)
class TicketAdjuntoAdmin(admin.ModelAdmin):
    list_display = ("ticket", "nombre", "subido_en", "subido_por")
    search_fields = ("ticket__numero_display", "ticket__numero", "nombre")
    readonly_fields = ("subido_en", "subido_por")


# ---------------------------------------------------------------------------
# Documentos y carpetas
# ---------------------------------------------------------------------------
class DocumentoInline(admin.TabularInline):
    model = Documento
    extra = 1
    fields = ("archivo", "nombre", "subido_en")
    readonly_fields = ("subido_en",)


@admin.register(DocumentoCarpeta)
class DocumentoCarpetaAdmin(admin.ModelAdmin):
    list_display = ("nombre", "slug", "seccion", "num_documentos", "creado_en")
    search_fields = ("nombre", "slug")
    prepopulated_fields = {}
    inlines = [DocumentoInline]

    @admin.display(description="Documentos")
    def num_documentos(self, obj):
        return obj.documentos.count()


@admin.register(Documento)
class DocumentoAdmin(admin.ModelAdmin):
    list_display = ("nombre", "carpeta", "subido_en")
    list_select_related = ("carpeta",)
    search_fields = ("nombre", "carpeta__nombre")
    list_filter = ("carpeta",)
    readonly_fields = ("subido_en",)


@admin.register(Seccion)
class SeccionAdmin(admin.ModelAdmin):
    list_display = ("nombre",)
    search_fields = ("nombre",)
